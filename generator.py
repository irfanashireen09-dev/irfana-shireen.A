"""Formatting utilities: sanitize text and export to HTML preview, DOCX and PDF."""
import html
import io
import os
import re

from ai_core.gemini_generator import split_terms
from config import DOC_LOGO_PATH, FOOTER_TEXT

_REPLACEMENTS = {
    "\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"',
    "\u2013": "-", "\u2014": "-", "\u2026": "...", "\u00a0": " ",
    "\u200b": "", "\u2022": "- ", "\u25cf": "- ", "\ufeff": "",
}


# ---------------------------------------------------------------- sanitize
def sanitize_text(text: str) -> str:
    """Remove typographic quotes / markdown noise so every export format stays clean."""
    if not text:
        return ""
    for old, new in _REPLACEMENTS.items():
        text = text.replace(old, new)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)            # bold
    text = re.sub(r"__([^_\n]+?)__", r"\1", text)           # bold (underscore form)
    text = re.sub(r"(?m)^\s*[-*]{3,}\s*$", "", text)        # horizontal rules
    text = re.sub(r"(?m)^(\s*)\*\s+", r"\1- ", text)        # "* item" -> "- item"
    text = re.sub(r"(?<!\*)\*(?!\s)([^*\n]+?)\*(?!\*)", r"\1", text)  # italics
    text = text.replace("`", "")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# ---------------------------------------------------------------- parsing
def _is_heading(line: str) -> bool:
    if len(line) > 80:
        return False
    if line.endswith(":") and ". " not in line:
        return True
    if re.match(r"^\d+(\.\d+)*[.)]\s+[A-Za-z]", line) and not line.endswith((".", ";", ",")):
        return True
    letters = re.sub(r"[^A-Za-z]", "", line)
    return len(letters) >= 4 and line.upper() == line and not line.endswith((".", ";", ","))


def parse_blocks(text: str):
    """Turn raw text into [(kind, text)] where kind is heading | bullet | para."""
    blocks = []
    for raw in sanitize_text(text).split("\n"):
        line = raw.strip()
        if not line:
            continue
        if line.startswith("#"):
            blocks.append(("heading", line.lstrip("#").strip()))
        elif re.match(r"^-\s+", line):
            blocks.append(("bullet", re.sub(r"^-\s+", "", line)))
        elif _is_heading(line) or (not blocks and len(line) <= 80 and not line.endswith((".", ":"))):
            blocks.append(("heading", line))  # section heading, or the document title on line 1
        else:
            blocks.append(("para", line))
    return blocks


# ---------------------------------------------------------------- HTML preview
def format_html_preview(text: str) -> str:
    """Return styled HTML (no newlines, so Streamlit markdown never breaks it)."""
    out, in_list = [], False
    for kind, content in parse_blocks(text):
        safe = html.escape(content)
        if kind != "bullet" and in_list:
            out.append("</ul>")
            in_list = False
        if kind == "heading":
            out.append(f"<h4 style='color:#8ab4ff;margin:18px 0 6px 0;'>{safe}</h4>")
        elif kind == "bullet":
            if not in_list:
                out.append("<ul style='margin:4px 0 8px 18px;'>")
                in_list = True
            out.append(f"<li style='margin-bottom:4px;'>{safe}</li>")
        else:
            out.append(f"<p style='margin:0 0 10px 0;line-height:1.55;text-align:justify;'>{safe}</p>")
    if in_list:
        out.append("</ul>")
    return "".join(out)


# ---------------------------------------------------------------- DOCX
def format_docx(text: str, doc_type: str, terms: str = "") -> bytes:
    """Build a Word document: logo, Times New Roman, terms table, footer."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    doc = Document()
    doc.core_properties.title = doc_type
    section = doc.sections[0]
    section.left_margin = section.right_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    if os.path.exists(DOC_LOGO_PATH):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(DOC_LOGO_PATH, width=Inches(2.4))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(doc_type.upper())
    run.bold = True
    run.font.size = Pt(18)
    title.paragraph_format.space_after = Pt(14)

    term_list = split_terms(terms)
    if term_list:
        h = doc.add_paragraph()
        h.add_run("Summary of Key Terms").bold = True
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.autofit = False
        table.rows[0].cells[0].text = "No."
        table.rows[0].cells[1].text = "Term / Condition"
        for cell in table.rows[0].cells:
            for r in cell.paragraphs[0].runs:
                r.bold = True
        for i, t in enumerate(term_list, start=1):
            row = table.add_row().cells
            row[0].text = str(i)
            row[1].text = t
        for row in table.rows:
            row.cells[0].width = Inches(0.6)
            row.cells[1].width = Inches(5.9)
        doc.add_paragraph()

    blocks = parse_blocks(text)
    if blocks and blocks[0][0] in ("heading", "para") and blocks[0][1].strip(" :").lower() == doc_type.strip().lower():
        blocks = blocks[1:]  # title already shown

    for kind, content in blocks:
        if kind == "heading":
            p = doc.add_paragraph()
            r = p.add_run(content)
            r.bold = True
            r.font.size = Pt(13)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.keep_with_next = True
        elif kind == "bullet":
            doc.add_paragraph(content, style="List Bullet")
        else:
            p = doc.add_paragraph(content)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(6)

    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run(FOOTER_TEXT)
    fr.italic = True
    fr.font.size = Pt(9)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- PDF
def _latin1(s: str) -> str:
    """Core PDF fonts only support latin-1."""
    return s.encode("latin-1", "replace").decode("latin-1")


def _make_pdf_class():
    from fpdf import FPDF

    class LegalPDF(FPDF):
        def __init__(self, title: str):
            super().__init__(format="A4")
            self.doc_title = _latin1(title)
            self.set_margins(20, 20, 20)
            self.set_auto_page_break(True, margin=25)

        def header(self):
            if os.path.exists(DOC_LOGO_PATH):
                self.image(DOC_LOGO_PATH, x=(self.w - 45) / 2, y=8, w=45)
            self.set_y(26)
            self.set_font("Helvetica", "B", 13)
            self.cell(0, 8, self.doc_title, align="C")
            self.ln(14)

        def footer(self):
            self.set_y(-18)
            self.set_font("Helvetica", "I", 8)
            self.cell(0, 5, _latin1(FOOTER_TEXT), align="C")
            self.ln(4)
            self.cell(0, 5, f"Page {self.page_no()}", align="C")

    return LegalPDF


def _write(pdf, height, txt, indent=0):
    pdf.set_x(pdf.l_margin + indent)
    pdf.multi_cell(0, height, _latin1(txt))
    pdf.set_x(pdf.l_margin)


def format_pdf(text: str, doc_type: str, terms: str = "") -> bytes:
    """Build a branded PDF with logo header, bold headings, bullet terms and footer."""
    pdf = _make_pdf_class()(doc_type)
    pdf.add_page()

    term_list = split_terms(terms)
    if term_list:
        pdf.set_font("Helvetica", "B", 11)
        _write(pdf, 7, "Summary of Key Terms")
        pdf.set_font("Helvetica", "", 11)
        for t in term_list:
            _write(pdf, 6, "- " + t, indent=4)
        pdf.ln(3)

    blocks = parse_blocks(text)
    if blocks and blocks[0][1].strip(" :").lower() == doc_type.strip().lower():
        blocks = blocks[1:]

    for kind, content in blocks:
        if kind == "heading":
            pdf.ln(2)
            pdf.set_font("Helvetica", "B", 11)
            _write(pdf, 7, content)
            pdf.set_font("Helvetica", "", 11)
        elif kind == "bullet":
            pdf.set_font("Helvetica", "", 11)
            _write(pdf, 6, "- " + content, indent=4)
        else:
            pdf.set_font("Helvetica", "", 11)
            _write(pdf, 6, content)
            pdf.ln(1.5)

    try:
        out = pdf.output(dest="S")
    except TypeError:
        out = pdf.output()
    return out.encode("latin-1") if isinstance(out, str) else bytes(out)